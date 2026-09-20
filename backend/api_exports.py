import io
import json
from datetime import datetime

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.styles import PatternFill

from core import db, require_admin, serialize

router = APIRouter(prefix="/export", tags=["export"])

ENTITY_COLLECTIONS = {
    "loans": "loans",
    "insurance": "insurance",
    "projects": "projects",
    "farms": "farms",
    "savings": "savings",
    "pf_ppf": "pf_ppf",
    "pf-ppf": "pf_ppf",
    "investments": "investments",
    "assets": "assets",
    "liabilities": "liabilities",
    "transactions": "transactions",
    "lendings": "lendings",
    "family": "family_members",
    "parties": "parties",
    "work_logs": "work_logs",
}


def _normalize_value(value):
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return str(value)


def build_workbook(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"

    if not rows:
        sheet.append(["No data"])
        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return buffer

    columns = []
    for row in rows:
        for key in row.keys():
            if key not in columns:
                columns.append(key)

    sheet.append(columns)
    for row in rows:
        sheet.append([_normalize_value(row.get(col, "")) for col in columns])

    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for column_cells in sheet.columns:
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 12), 45)

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


async def _fetch_rows(entity: str, query_filters: dict):
    normalized = entity.strip().lower().replace("-", "_")
    collection_name = ENTITY_COLLECTIONS.get(normalized)
    if collection_name is None:
        raise HTTPException(status_code=400, detail=f"Unsupported export entity: {entity}")

    collection = db[collection_name]
    query = {"deleted_at": {"$exists": False}}
    for key, value in query_filters.items():
        if key in {"entity", "format", "download"}:
            continue
        if value in {None, ""}:
            continue
        query[key] = value

    docs = await collection.find(query).sort([("created_at", -1)]).to_list(50000)
    return [serialize(doc) for doc in docs]


@router.get("")
async def export_data(request: Request, user: dict = Depends(require_admin), entity: str = ""):
    if not entity:
        raise HTTPException(status_code=400, detail="Missing export entity")

    query_filters = {}
    for key, value in request.query_params.items():
        if key in {"entity", "format", "download"}:
            continue
        query_filters[key] = value

    rows = await _fetch_rows(entity, query_filters)
    workbook = build_workbook(rows)
    safe_entity = entity.strip().lower().replace("-", "_")
    filename = f"{safe_entity}_export_{__import__('datetime').datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"

    return StreamingResponse(
        workbook,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/financial-snapshot")
async def financial_snapshot(user: dict = Depends(require_admin)):
    """Portable, self-attested household snapshot. It is clearly labelled as a
    user-maintained report, not bank or advisor verification."""
    from api_wealth import networth_data
    from api_planning import planning_overview, ownership_summary
    wealth = await networth_data()
    plan = await planning_overview(user)
    ownership = await ownership_summary(user)
    wb = Workbook(); summary = wb.active; summary.title = "Financial snapshot"
    rows = [["Nivara Financial Snapshot"], ["Generated at", datetime.utcnow().isoformat() + "Z"], ["Status", "Self-attested household record — verify with original documents before sharing"], [], ["Net worth", wealth["net_worth"]], ["Total assets", wealth["total_assets"]], ["Total liabilities", wealth["total_liabilities"]], ["Cash on hand", plan["forecast"]["cash_on_hand"]], ["Minimum projected balance", plan["forecast"]["minimum_balance"]], ["Minimum balance date", plan["forecast"]["minimum_balance_date"]], ["Financial health score", plan["health"]["score"]]]
    for row in rows: summary.append(row)
    summary["A1"].font = Font(bold=True, size=16); summary["A1"].fill = PatternFill("solid", fgColor="D1FAE5")
    for c in summary[1]: c.font = Font(bold=True)
    summary.column_dimensions["A"].width = 32; summary.column_dimensions["B"].width = 70
    own = wb.create_sheet("Ownership")
    own.append(["Owner / entity", "Tracked net value"])
    for row in ownership["owners"]: own.append([row["name"], row["value"]])
    own.append([]); own.append(["My finances", ownership["views"]["mine"]]); own.append(["Family finances", ownership["views"]["family"]]); own.append(["Farm / business finances", ownership["views"]["farm_business"]]); own.append(["Consolidated", ownership["views"]["consolidated"]])
    commits = wb.create_sheet("Commitments")
    commits.append(["Due date", "Action", "Detail", "Amount", "Type"])
    for event in plan["calendar"]: commits.append([event["due_date"], event["title"], event["detail"], event["amount"], event["kind"]])
    docs = wb.create_sheet("Documents")
    docs.append(["File", "Category", "Linked record type", "Linked record id", "Uploaded"])
    for doc in await db.documents.find({"deleted_at": {"$exists": False}}).sort("created_at", -1).to_list(5000): docs.append([doc.get("filename"), doc.get("category"), doc.get("related_entity_type"), doc.get("related_entity_id"), str(doc.get("created_at", ""))])
    for sheet in wb.worksheets:
        for cell in sheet[1]: cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(max(max(len(str(c.value or "")) for c in column) + 2, 14), 55)
    buffer = io.BytesIO(); wb.save(buffer); buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="nivara_financial_snapshot_{datetime.utcnow().strftime("%Y%m%d")}.xlsx"'})
