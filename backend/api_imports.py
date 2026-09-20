"""Provider-neutral, review-first financial import pipeline."""
from __future__ import annotations
import csv, hashlib, io, os, re
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from core import _client, db, now_utc, require_admin, round2
from financial_services import entity_for_investment, fingerprint, normalized_key, record_imported_transaction

router = APIRouter(tags=["imports"])
MAX_BYTES, SUPPORTED = 10 * 1024 * 1024, {".csv", ".xlsx", ".xlsm"}

class CommitImportRequest(BaseModel):
    row_numbers: list[int] = Field(min_length=1)
    owner: str = Field(default="Self", min_length=1, max_length=120)
    occurred_on: str | None = None

def _number(value):
    try: return round2(float(str(value or "0").replace(",", "").replace("₹", "").strip()))
    except (ValueError, TypeError): return 0.0

def _clean_row(row):
    return {str(key or "").strip().lower().replace(" ", "_"): value for key, value in row.items() if str(key or "").strip()}

def detect_provider(filename, headers):
    signal = f"{filename} {' '.join(headers)}".lower()
    if "zerodha" in signal or "tradingsymbol" in signal or "isin" in signal: return "ZERODHA", "BROKER_EXPORT", .94
    if "groww" in signal: return "GROWW", "BROKER_EXPORT", .90
    if "upstox" in signal: return "UPSTOX", "BROKER_EXPORT", .90
    return "GENERIC", "TABULAR_FINANCIAL_EXPORT", .58

def classify_row(row):
    clean = _clean_row(row); text = " ".join(str(value or "") for value in clean.values()).lower()
    name = clean.get("symbol") or clean.get("tradingsymbol") or clean.get("instrument") or clean.get("scrip") or clean.get("name") or "Imported investment"
    pnl = next((_number(clean.get(key)) for key in ("realized_pnl", "realised_pnl", "p&l", "pnl", "profit_loss") if clean.get(key) not in (None, "")), None)
    quantity = next((_number(clean.get(key)) for key in ("quantity", "qty", "net_quantity")), 0)
    value = next((_number(clean.get(key)) for key in ("current_value", "market_value", "value", "closing_value", "amount")), 0)
    cost = next((_number(clean.get(key)) for key in ("invested", "cost", "buy_value", "average_price")), 0)
    isin = str(clean.get("isin") or "").strip()
    if pnl is not None or "p&l" in text or "profit" in text or "loss" in text:
        return {"kind":"HISTORICAL_PNL", "confidence":.91, "name":str(name), "isin":isin, "amount":pnl or 0, "reason":"Recognised a realised P&L field; it will not be treated as a current holding."}
    if name and (quantity or value or cost):
        return {"kind":"HOLDING", "confidence":.86, "name":str(name), "isin":isin, "quantity":quantity, "current_value":value or cost, "cost":cost or value, "reason":"Recognised an instrument with holding, quantity, cost or valuation fields."}
    return {"kind":"NEEDS_GUIDANCE", "confidence":.35, "name":str(name), "isin":isin, "reason":"No safe financial interpretation was found. Choose a mapping in Review Inbox."}

def parse_rows(filename, content):
    ext = os.path.splitext(filename.lower())[1]
    if ext == ".csv": return list(csv.DictReader(io.StringIO(content.decode("utf-8-sig", errors="replace"))))
    if ext in {".xlsx", ".xlsm"}:
        values = list(load_workbook(io.BytesIO(content), read_only=True, data_only=True).active.iter_rows(values_only=True))
        if not values: return []
        headers = [str(value or "").strip() for value in values[0]]
        return [dict(zip(headers, row)) for row in values[1:] if any(value not in (None, "") for value in row)]
    raise HTTPException(status_code=415, detail="This file type needs a parser adapter. CSV, XLSX and XLSM are currently supported for safe extraction.")

@router.post("/imports/analyze")
async def analyze_import(file: UploadFile = File(...), user: dict = Depends(require_admin)):
    filename, content = file.filename or "upload", await file.read()
    if not content: raise HTTPException(status_code=400, detail="The uploaded file is empty")
    if len(content) > MAX_BYTES: raise HTTPException(status_code=413, detail="Import file is too large (10MB maximum)")
    content_sha256 = hashlib.sha256(content).hexdigest(); previous = await db.import_runs.find_one({"content_sha256": content_sha256})
    if previous: return {"status":"ALREADY_IMPORTED", "run_id":str(previous["_id"]), "filename":previous.get("filename"), "message":"This exact document was already analyzed. Open its audit record instead of importing it again."}
    rows = parse_rows(filename, content)
    if not rows: raise HTTPException(status_code=400, detail="No tabular records were found")
    provider, document_type, confidence = detect_provider(filename, list(_clean_row(rows[0]).keys())); now = now_utc()
    run = {"filename":filename,"content_sha256":content_sha256,"provider":provider,"document_type":document_type,"document_confidence":confidence,"status":"REVIEW","created_at":now,"created_by":user["email"],"row_count":len(rows),"parser_version":"2026.09.20"}
    result = await db.import_runs.insert_one(run); run_id = str(result.inserted_id); candidates=[]
    for number, source in enumerate(rows[:3000], start=2):
        candidate = classify_row(source); row_fp = fingerprint(provider, candidate.get("isin") or candidate["name"], candidate["kind"], candidate.get("quantity", candidate.get("amount", 0)), candidate.get("current_value", candidate.get("amount", 0)))
        duplicate = await db.import_rows.find_one({"fingerprint":row_fp,"status":"COMMITTED"})
        candidate.update({"row_number":number,"fingerprint":row_fp,"duplicate":bool(duplicate),"source":source})
        await db.import_rows.insert_one({"run_id":run_id,"row_number":number,"fingerprint":row_fp,"candidate":candidate,"source_row":source,"status":"DUPLICATE" if duplicate else "REVIEW","created_at":now})
        candidates.append(candidate)
    return {"status":"REVIEW","run_id":run_id,"filename":filename,"provider":provider,"document_type":document_type,"document_confidence":confidence,"total_rows":len(rows),"candidates":candidates,"notice":"Nothing has changed in your finances. Review and approve only recognised, non-duplicate rows."}

async def _commit_row(row, run, request, session=None):
    candidate, kind = row["candidate"], row["candidate"]["kind"]
    if kind == "NEEDS_GUIDANCE": return "NEEDS_GUIDANCE"
    if row.get("status") == "DUPLICATE" or candidate.get("duplicate"): return "DUPLICATE"
    owner = request.owner.strip(); occurrence = request.occurred_on or now_utc().date().isoformat()
    entity = await entity_for_investment(candidate["name"], owner, "Equity", str(run["_id"]), session=session)
    if kind == "HOLDING":
        identity = normalized_key(candidate.get("isin") or candidate["name"], owner, "Equity"); existing = await db.investments.find_one({"identity_key":identity,"deleted_at":{"$exists":False}}, session=session)
        update = {"name":candidate["name"],"symbol":candidate["name"],"isin":candidate.get("isin"),"type":"Stocks","asset_class":"Equity","cost":candidate.get("cost",0),"current_value":candidate.get("current_value",0),"quantity":candidate.get("quantity",0),"owner":owner,"ownership_percent":100,"identity_key":identity,"canonical_entity_id":str(entity["_id"]),"source_import_id":str(run["_id"]),"source_document":run["filename"],"source_row":row["row_number"],"updated_at":now_utc()}
        if existing: await db.investments.update_one({"_id":existing["_id"]},{"$set":update}, session=session); outcome="CONSOLIDATED"
        else: update["created_at"]=now_utc(); await db.investments.insert_one(update, session=session); outcome="CREATED"
        await db.holdings.update_one({"identity_key":identity},{"$set":{"entity_id":str(entity["_id"]),"identity_key":identity,"symbol":candidate["name"],"isin":candidate.get("isin"),"quantity":candidate.get("quantity",0),"cost":candidate.get("cost",0),"market_value":candidate.get("current_value",0),"valuation_date":occurrence,"source_import_id":str(run["_id"]),"source_row":row["row_number"],"updated_at":now_utc()},"$setOnInsert":{"created_at":now_utc()}}, upsert=True, session=session)
        await record_imported_transaction(run_id=str(run["_id"]),row_number=row["row_number"],entity_id=str(entity["_id"]),kind="HOLDING_VALUATION",amount=candidate.get("current_value",0),occurred_on=occurrence,source_reference=run["filename"], session=session)
    else:
        amount=candidate.get("amount",0)
        await record_imported_transaction(run_id=str(run["_id"]),row_number=row["row_number"],entity_id=str(entity["_id"]),kind="HISTORICAL_PNL",amount=amount,occurred_on=occurrence,source_reference=run["filename"], session=session)
        await db.transactions.insert_one({"type":"INCOME" if amount>=0 else "EXPENSE","amount":abs(amount),"category":"Investment P&L","source":run["provider"],"title":candidate["name"],"description":"Reviewed historical broker P&L import","date":occurrence,"scope":"PERSONAL","source_import_id":str(run["_id"]),"source_document":run["filename"],"source_row":row["row_number"],"created_at":now_utc(),"created_by":run["created_by"]}, session=session); outcome="CREATED"
    await db.import_rows.update_one({"_id":row["_id"]},{"$set":{"status":"COMMITTED","committed_at":now_utc(),"outcome":outcome}}, session=session)
    return outcome

@router.post("/imports/{run_id}/commit")
async def commit_import(run_id: str, request: CommitImportRequest, user: dict = Depends(require_admin)):
    from bson import ObjectId
    try: object_id=ObjectId(run_id)
    except Exception as exc: raise HTTPException(status_code=400, detail="Invalid import run") from exc
    run=await db.import_runs.find_one({"_id":object_id,"status":"REVIEW"})
    if not run: raise HTTPException(status_code=404, detail="Import review was not found or has already been committed")
    selected=set(request.row_numbers); rows=await db.import_rows.find({"run_id":run_id,"row_number":{"$in":list(selected)}}).to_list(3000)
    if len(rows)!=len(selected): raise HTTPException(status_code=400, detail="One or more selected rows do not belong to this import")
    summary={"created":0,"consolidated":0,"duplicates":0,"needs_guidance":0}; keys={"CREATED":"created","CONSOLIDATED":"consolidated","DUPLICATE":"duplicates","NEEDS_GUIDANCE":"needs_guidance"}
    try:
        async with await _client.start_session() as session:
            async with session.start_transaction():
                for row in rows: summary[keys[await _commit_row(row,run,request,session)]]+=1
    except Exception as exc:
        if "transaction numbers are only allowed" not in str(exc).lower() and "replica set" not in str(exc).lower(): raise
        for row in rows: summary[keys[await _commit_row(row,run,request)]]+=1
    await db.import_runs.update_one({"_id":object_id},{"$set":{"status":"COMMITTED","committed_at":now_utc(),"committed_by":user["email"],"summary":summary}})
    return {"run_id":run_id,"summary":summary}

@router.get("/imports")
async def list_imports(user: dict = Depends(require_admin)):
    rows=await db.import_runs.find({}).sort("created_at",-1).to_list(200)
    return [{"id":str(row["_id"]),"filename":row.get("filename"),"provider":row.get("provider"),"document_type":row.get("document_type"),"status":row.get("status"),"row_count":row.get("row_count"),"created_at":row.get("created_at"),"summary":row.get("summary",{})} for row in rows]
